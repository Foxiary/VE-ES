"""
applyui.py - write a sheet's interface rows back into the SYSTEM.cpk databases.

    python tools/applyui.py work/glossary.xlsx work/out
    python tools/applyui.py work/virche_vi.xlsx work/out --dry-run
    python tools/applyui.py work/glossary.xlsx work/out --cpk dist/SYSTEM_text.cpk

The counterpart to `applyvi.py`, for the other half of the text. `applyvi.py`
patches `.DAT` blocks in STORY.cpk and its id regex refuses a `ui` row on
purpose; this one takes exactly those rows and rebuilds the `.gbin` / `.gstr`
files they came from.

WHERE IT SITS IN THE BUILD
    It writes rebuilt database files into a directory, which is step 3 of
    `build.py` - the same slot `translate_glossary.py` fills, and the same
    directory `cpk.py repack` then folds into the CPK beside the fonts. So the
    normal route is `applyui.py sheet.xlsx work/out` followed by the usual
    build; `--cpk` is a shortcut that repacks the source archive with whatever
    is in the output directory, for checking one screen without a full rebuild.

    `--src` defaults to SYSTEM.cpk rather than `work/stock`, because the stock
    directory only holds the four files named in `fonts.json` and a sheet can
    carry rows from any database - the Glossary alone also touches strGame.gstr,
    which is not one of them. Reading the archive means a sheet is never
    silently short a file.

ADDRESSED BY ID, AND THE ADDRESS IS CHECKED
    `translate_glossary.py` matches by content - it looks for the English string
    anywhere in the pool. That is the approach CLAUDE.md warns about: two
    records can hold the same text for different reasons, and a near miss
    silently replaces nothing. A mksheet id is an address,
    `record.column___pooloffset_role`, so a row names its cell outright. All
    three parts are proven before anything is written: the cell must exist, it
    must still point at that pool offset, and the string there must be the
    sheet's source text. A sheet built from a different build fails the second
    or third check and is reported, not guessed at.

ONE ROW UPDATES EVERY CELL HOLDING THAT STRING
    `gbnl.build()` is keyed by text, not by cell - it rewrites every cell
    pointing at a string it is given. mksheet emits one row per distinct string
    for that reason, and it is also why two rows in one file that share a source
    but disagree on the translation are refused: whichever won would silently
    decide the other. In dbDictionary that is load-bearing rather than an edge
    case, since columns 8 and 16 of every record point at the same entry.

THE POOL IS REBUILT, SO LENGTH DOES NOT MATTER
    Unlike a `.DAT` block there is no capacity to fit inside: `gbnl.build()`
    lays out a fresh string pool and remaps every offset, so Vietnamese is free
    to run longer than English. What the file cannot survive is a lost schema
    block, which is why the result is parsed back and every cell compared before
    anything reaches disk - see `verify()`.
"""
import argparse
import collections
import os
import re
import struct
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gbnl import GBNL                                # noqa: E402
from cpk import CPK                                  # noqa: E402
from linebreak import to_game, canon                 # noqa: E402
# Borrowed rather than copied: the header lookup has to stay identical to the
# story path, or the two tools disagree about which column a sheet keeps its
# translation in.
from applyvi import sheet_columns, cell_str         # noqa: E402

# Only the roles a database column can carry. A story row pasted into one of
# these sheets names an offset into a .DAT and is refused rather than applied.
ID_RX = re.compile(r'^(\d+)\.(\d+)___([0-9A-Fa-f]+)_(ui|name|title)$')

# Commands that carry an argument, and so mean something specific. These must
# survive a translation exactly: `#NAME[1]` is the name the player chose,
# `#Color[8]`/`#Color[0]` open and close a coloured run, and the ten `#PosX[%d]`
# of IDS_DICTIONARY_TAB are the Glossary's tab stops.
COMMAND_RX = re.compile(r'#(?:NAME|Color|PosX)\[[^\]]*\]')

# `#n` is deliberately NOT in there. It is a line break, and a translator
# re-breaking a paragraph to fit Vietnamese is doing their job, not losing a
# command: a database string goes back through `gbnl.build()`, which rebuilds
# the pool, so neither its length nor its line count is constrained. Requiring
# the counts to match refused 78 finished rows on this project's first real
# sheet. The delta is counted and reported instead.
BREAK = '#n'


def commands(text):
    return collections.Counter(COMMAND_RX.findall(text or ''))


SUFFIXES = ('.gbin', '.gstr')


def sources(path):
    """{basename: bytes} from a .cpk or from a directory of unpacked files."""
    if os.path.isdir(path):
        return {n: open(os.path.join(path, n), 'rb').read()
                for n in sorted(os.listdir(path)) if n.endswith(SUFFIXES)}
    c = CPK(path)
    return {os.path.basename(full): c.read(row)
            for _i, full, row in c.files() if full.endswith(SUFFIXES)}


def cell_offset(g, record, col):
    """The pool offset a string cell holds, or None if it is not a string cell."""
    if record >= g.n or col not in g.str_cols:
        return None
    return struct.unpack_from('<Q', g.raw, g.tbl + record * g.stride + col)[0]


def plan(ws, g):
    """({old text: new text}, tally, [refusal lines]) for one database file."""
    rows = ws.iter_rows(values_only=True)
    c_id, c_src, c_tgt = sheet_columns(next(rows, ()))
    repl, st, notes = {}, collections.Counter(), []

    for r in rows:
        if not r or c_id >= len(r) or not r[c_id]:
            continue
        rid = str(r[c_id]).strip()
        m = ID_RX.match(rid)
        if not m:
            st['id khong doc duoc'] += 1
            continue
        st['rows'] += 1
        record, col = int(m.group(1)), int(m.group(2)) * 8
        rel, role = int(m.group(3), 16), m.group(4)
        # Stripped first, then converted - a trailing newline is invisible in
        # Excel and would otherwise come back as a trailing #n.
        src = to_game(cell_str(r[c_src]).strip()) if c_src < len(r) else ''
        tgt = to_game(cell_str(r[c_tgt]).strip()) if c_tgt < len(r) else ''

        if not tgt:
            st['chua dich'] += 1
            continue
        if cell_offset(g, record, col) != rel:
            st['id khong tro vao o chuoi'] += 1
            notes.append('%s: o (%d, %d) khong con giu offset %#x'
                         % (rid, record, col, rel))
            continue
        # canon(): the sheet's source column may carry padding around its
        # breaks from whatever edited it. The id already names the exact cell;
        # this only has to confirm the sheet belongs to this build.
        if canon(g.text(rel)) != canon(src):
            st['van ban khong khop'] += 1        # sheet built from another dump
            continue
        if commands(tgt) != commands(src):
            st['mat lenh'] += 1
            notes.append('%s: lenh %s -> %s'
                         % (rid, dict(commands(src)), dict(commands(tgt))))
            continue
        if tgt.count(BREAK) != src.count(BREAK):
            st['ngat dong lai'] += 1            # allowed; counted so it is visible
        # A flowchart title is "<japanese key>@<display text>" and the key is
        # what the scene table is looked up by. Losing it corrupts the lookup
        # rather than just the text, so the row is refused, never repaired.
        if role == 'title' and '@' in src \
                and not tgt.startswith(src.split('@', 1)[0] + '@'):
            st['title mat khoa'] += 1
            notes.append('%s: mat khoa %r' % (rid, src.split('@', 1)[0]))
            continue
        # KEY ON THE POOL'S OWN TEXT, NOT ON THE SHEET'S SOURCE COLUMN.
        # `gbnl.build()` looks each pool string up in this dict verbatim, so a
        # key that differs from it by even one space replaces nothing - and the
        # row still counts as applied, which is the worst kind of wrong. The
        # comparison above is deliberately tolerant of editor damage; the key
        # cannot be. This shipped once with every multi-line glossary entry
        # silently left in English while the single-line names translated fine.
        key = g.text(rel)
        if tgt == key:
            st['giong ban goc'] += 1
            continue
        if key in repl and repl[key] != tgt:
            # build() is keyed by text: whichever won would decide the other.
            st['trung nguon khac ban dich'] += 1
            notes.append('%s: %r co hai ban dich khac nhau' % (rid, key[:40]))
            continue
        repl[key] = tgt
        st['applied'] += 1

    return repl, st, notes


def verify(before, data, repl):
    """Parse the rebuilt file back and prove every string cell reads as intended.

    A clean `build()` is not evidence on its own - the schema block sits between
    the record table and the pool, and a file that lost it still parses here
    while crashing the game the moment a screen reads that table. So the output
    is checked three ways: the table's shape is unchanged, the schema block is
    byte-identical, and every cell holds either its translation or exactly what
    it held before.
    """
    after = GBNL(data)
    if (after.n, after.stride, after.ncol, after.str_cols) != \
            (before.n, before.stride, before.ncol, before.str_cols):
        return 'shape doi khac'
    if data[after.tbl_end:after.pool] != before.raw[before.tbl_end:before.pool]:
        return 'khoi schema doi khac'
    for i in range(before.n):
        for col in before.str_cols:
            old = before.text(cell_offset(before, i, col))
            want = repl.get(old, old)
            got = after.text(cell_offset(after, i, col))
            if got != want:
                return 'o (%d, %d): doc ra %r, can %r' % (i, col, got[:40], want[:40])
    return None


def main():
    ap = argparse.ArgumentParser(
        description='write a sheet\'s interface rows into the SYSTEM databases')
    ap.add_argument('xlsx', help='a mksheet.py / glossary.py workbook')
    ap.add_argument('out_dir', help='where the rebuilt .gbin/.gstr are written')
    ap.add_argument('--src', default='SYSTEM.cpk',
                    help='SYSTEM.cpk or a directory of unpacked databases')
    ap.add_argument('--sheet', action='append', help='limit to these worksheets')
    ap.add_argument('--cpk', help='also repack --src with everything in out_dir')
    ap.add_argument('--dry-run', action='store_true')
    a = ap.parse_args()

    import openpyxl
    wb = openpyxl.load_workbook(a.xlsx, read_only=True, data_only=True)
    src = sources(a.src)
    stem = {os.path.splitext(n)[0]: n for n in src}

    names = [n for n in wb.sheetnames if n in stem]
    if a.sheet:
        names = [n for n in names if n in a.sheet]
    if not names:
        sys.exit('khong sheet nao trung ten file trong %s' % a.src)
    other = [n for n in wb.sheetnames if n not in stem and n != '_README']
    if other:
        print('bo qua %d sheet khong phai database: %s%s'
              % (len(other), ', '.join(other[:6]), ' ...' if len(other) > 6 else ''))

    if not a.dry_run:
        os.makedirs(a.out_dir, exist_ok=True)

    tot = collections.Counter()
    written = []
    print('\n%-24s %7s %8s %8s %16s' % ('file', 'dong', 'ap dung', 'bo qua', 'byte'))
    for name in names:
        base = stem[name]
        g = GBNL(src[base])
        if g.build({}) != g.raw:
            print('!! %s: round-trip khong sach, bo qua' % base)
            tot['file hong'] += 1
            continue

        repl, st, notes = plan(wb[name], g)
        tot.update(st)
        left = st['rows'] - st['applied']
        if not repl:
            print('%-24s %7d %8d %8d %16s' % (base, st['rows'], 0, left, '-'))
            continue

        data = g.build(repl)
        bad = verify(g, data, repl)
        if bad:
            print('!! %s: KIEM TRA HONG - %s' % (base, bad))
            tot['file hong'] += 1
            continue

        print('%-24s %7d %8d %8d %16s' %
              (base, st['rows'], st['applied'], left,
               '%s -> %s' % (format(len(g.raw), ','), format(len(data), ','))))
        for note in notes[:4]:
            print('    !! %s' % note)
        if len(notes) > 4:
            print('    !! ... va %d dong nua bi tu choi' % (len(notes) - 4))
        if not a.dry_run:
            with open(os.path.join(a.out_dir, base), 'wb') as fh:
                fh.write(data)
        written.append(base)

    print('-' * 68)
    print('ap dung %s dong tren %d file%s'
          % (format(tot['applied'], ','), len(written),
             ' (dry-run, khong ghi gi)' if a.dry_run else ''))
    if tot['ngat dong lai']:
        print('  (%d dong da ap dung nhung nguoi dich ngat dong khac ban goc)'
              % tot['ngat dong lai'])
    for k in ('chua dich', 'giong ban goc', 'mat lenh',
              'van ban khong khop', 'title mat khoa', 'trung nguon khac ban dich',
              'id khong tro vao o chuoi', 'id khong doc duoc', 'file hong'):
        if tot[k]:
            print('  bo qua - %-26s %s' % (k, format(tot[k], ',')))

    if a.cpk and not a.dry_run:
        if os.path.isdir(a.src):
            sys.exit('--cpk can --src la mot .cpk')
        c = CPK(a.src)
        have = {n.lower(): os.path.join(a.out_dir, n) for n in os.listdir(a.out_dir)}
        rep = {full: open(have[os.path.basename(full).lower()], 'rb').read()
               for _i, full, _r in c.files()
               if os.path.basename(full).lower() in have}
        print('\nrepack: %d file -> %s' % (len(rep), a.cpk))
        c.repack(a.cpk, rep)

    return 1 if tot['file hong'] else 0


if __name__ == '__main__':
    sys.exit(main())
