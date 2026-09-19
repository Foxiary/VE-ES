"""
checksheet.py - check a translation spreadsheet against the game's script files.

    python checksheet.py <sheet.xlsx> <script_dir> [--sheet 100]

The spreadsheet carries one worksheet per script file (named 100, 101, ...) with
columns: ID | <source text> | Vietnamese | Note. Each ID embeds the byte offset
of the string inside that .DAT, e.g.

    81___4E24C_name   -> offset 0x4E24C, a speaker name
    1___48EFC_text    -> offset 0x48EFC, a dialogue line

Those offsets are only valid for the exact build the sheet was made from. This
tool reports, per worksheet, how many offsets land on a string that matches the
source column - which is the quickest way to tell whether a sheet belongs to the
build you are holding.

A healthy result is a match rate near 100%. A low rate means the sheet was made
from a different region or version, and applying it would corrupt the script.
"""
import argparse
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from linebreak import to_game                    # noqa: E402

ID_RX = re.compile(r'^([\d.]+)_+([0-9A-Fa-f]+)_(text|name|choice|title|var|ui)$')


def load_sheet(path):
    import openpyxl
    return openpyxl.load_workbook(path, data_only=True)


def rows_of(ws):
    """Yield (id, source, target) for rows with a parseable ID."""
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        m = ID_RX.match(str(r[0]).strip())
        if not m:
            continue
        src = (r[1] or '') if len(r) > 1 else ''
        tgt = (r[2] or '') if len(r) > 2 else ''
        # The sheet holds real newlines; the script holds #n. Both columns
        # are handed back in the game's spelling so callers can compare them
        # against a block byte for byte.
        yield (str(r[0]).strip(), int(m.group(2), 16), m.group(3),
               to_game(str(src)), to_game(str(tgt)))


def cstr(d, off):
    e = d.find(b'\x00', off)
    if e < 0:
        return None
    try:
        return d[off:e].decode('utf-8')
    except UnicodeDecodeError:
        return None


def check_one(ws, dat):
    d = open(dat, 'rb').read()
    hit = miss = oob = blank = 0
    translated = 0
    for _rid, off, _kind, src, tgt in rows_of(ws):
        if tgt.strip():
            translated += 1
        if off >= len(d):
            oob += 1
            continue
        s = cstr(d, off)
        if s is None:
            miss += 1
        elif not src.strip():
            blank += 1                      # nothing to compare against
        elif s.strip() == src.strip():
            hit += 1
        else:
            miss += 1
    total = hit + miss + oob + blank
    rate = (hit / (hit + miss) * 100) if (hit + miss) else 0.0
    return dict(total=total, hit=hit, miss=miss, oob=oob, blank=blank,
                rate=rate, translated=translated, size=len(d))


def main():
    ap = argparse.ArgumentParser(description='check a translation sheet against script files')
    ap.add_argument('xlsx')
    ap.add_argument('script_dir', help='directory holding the .DAT files')
    ap.add_argument('--sheet', action='append', help='only check these worksheets')
    a = ap.parse_args()

    wb = load_sheet(a.xlsx)
    names = a.sheet or [n for n in wb.sheetnames if n.isdigit()]
    print('%-7s %8s %7s %7s %6s %6s %8s' %
          ('sheet', 'rows', 'match', 'miss', 'oob', 'dich', 'rate'))
    tot_hit = tot_miss = tot_oob = tot_tr = 0
    missing_files = []
    for n in names:
        dat = os.path.join(a.script_dir, '%s.DAT' % n)
        if not os.path.exists(dat):
            missing_files.append(n)
            continue
        r = check_one(wb[n], dat)
        tot_hit += r['hit']; tot_miss += r['miss']; tot_oob += r['oob']
        tot_tr += r['translated']
        flag = '' if r['rate'] > 95 else '  <-- lech'
        print('%-7s %8d %7d %7d %6d %6d %7.1f%%%s' %
              (n, r['total'], r['hit'], r['miss'], r['oob'],
               r['translated'], r['rate'], flag))
    denom = tot_hit + tot_miss
    print('-' * 56)
    print('tong: match %d, miss %d, ngoai file %d -> %.1f%%' %
          (tot_hit, tot_miss, tot_oob, (tot_hit / denom * 100) if denom else 0))
    print('so dong da dich: %d' % tot_tr)
    if missing_files:
        print('khong tim thay .DAT cho sheet: %s' % ', '.join(missing_files))
    if denom and tot_hit / denom > 0.95:
        print('\n=> Bang dich KHOP build nay.')
    else:
        print('\n=> Bang dich KHONG khop build nay (sai region hoac sai version).')


if __name__ == '__main__':
    sys.exit(main())
